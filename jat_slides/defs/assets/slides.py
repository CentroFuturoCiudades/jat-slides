import datetime
from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING, Literal, cast, overload
from zoneinfo import ZoneInfo

import numpy as np
import pandas as pd
from babel.dates import format_date
from cfc_dagster_utils.partitions import zone_partitions
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.parts.image import Image
from pptx.presentation import Presentation as PresentationType
from pptx.shapes.autoshape import Shape
from pptx.shapes.placeholder import PicturePlaceholder, TablePlaceholder
from pptx.shapes.shapetree import SlideShapes
from pptx.slide import SlideLayout
from pptx.util import Cm, Pt

import dagster as dg
from jat_slides.defs.partitions import mun_partitions

if TYPE_CHECKING:
    from pptx.text.text import TextFrame

RGB_BLUE = RGBColor(0x00, 0x70, 0xC0)
RGB_RED = RGBColor(0xFF, 0x00, 0x00)
TABLE_STYLE_ID = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"


@dataclass(frozen=True)
class Comparison:
    start_year: int
    lost_pop_fraction: float
    built_fraction: float
    population_map: Path


@dataclass(frozen=True)
class RadialComparison:
    start_year: int
    picture_path: Path
    table: pd.DataFrame


def find_layouts(pres: PresentationType) -> dict:
    layouts = {}
    for layout in pres.slide_layouts:
        name = layout.name
        if name == "Title Slide":
            layouts["title"] = layout
        elif name == "Picture with Title":
            layouts["picture_with_title"] = layout
        elif name == "Picture with Title and Content":
            layouts["picture_with_title_and_content"] = layout
        elif name == "Picture with Title, Content and Table":
            layouts["picture_with_title_content_table"] = layout
        elif name == "Section Divider":
            layouts["section"] = layout
    return layouts


@overload
def find_shape(shapes: SlideShapes, prefix: str, kind: Literal["shape"]) -> Shape: ...
@overload
def find_shape(
    shapes: SlideShapes, prefix: str, kind: Literal["picture"]
) -> PicturePlaceholder: ...
@overload
def find_shape(
    shapes: SlideShapes, prefix: str, kind: Literal["table"]
) -> TablePlaceholder: ...


def find_shape(
    shapes: SlideShapes, prefix: str, kind: Literal["shape", "picture", "table"]
) -> Shape | PicturePlaceholder | TablePlaceholder:
    type_map = {
        "shape": Shape,
        "picture": PicturePlaceholder,
        "table": TablePlaceholder,
    }
    for shape in shapes:
        if shape.name.startswith(prefix):
            if not isinstance(shape, type_map[kind]):
                err = (
                    "Shape with prefix '{prefix}' is not of type "
                    f"{type_map[kind].__name__}."
                )
                raise TypeError(err)
            return shape

    err = f"Shape with prefix '{prefix}' not found in the slide shapes."
    raise ValueError(err)


def add_title_slide(pres: PresentationType, title_layout: SlideLayout) -> None:
    title = pres.slides.add_slide(title_layout)

    placeholders = title.shapes
    title_shape = placeholders[0]
    date_shape = placeholders[2]

    if not isinstance(title_shape, Shape) or not isinstance(date_shape, Shape):
        err = "Placeholders are not of type Shape."
        raise TypeError(err)

    title_shape.text = "Dinámicas Urbanas"
    date_shape.text = format_date(
        datetime.datetime.now(tz=ZoneInfo("America/Mexico_City")).date(),
        locale="es",
        format="long",
    )


def add_section_slide(
    pres: PresentationType,
    section_layout: SlideLayout,
    section_name: str,
) -> None:
    section = pres.slides.add_slide(section_layout)
    section_title_shape = section.shapes[0]

    if not isinstance(section_title_shape, Shape):
        err = "Section title shape is not of type Shape."
        raise TypeError(err)
    section_title_shape.text = section_name


def add_highlighted_text_to_frame(
    shape: Shape,
    *,
    text_left: str,
    highlight: str,
    text_right: str,
    color: RGBColor,
) -> None:
    frame: TextFrame = shape.text_frame
    frame.clear()

    p = frame.paragraphs[0]
    run = p.add_run()
    run.text = text_left

    run = p.add_run()
    run.text = highlight
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = color

    run = p.add_run()
    run.text = text_right


def add_normal_text_to_shape(shape: Shape, text: str) -> None:
    frame: TextFrame = shape.text_frame
    frame.clear()

    p = frame.paragraphs[0]
    run = p.add_run()
    run.text = text


def add_picture_with_highlight_slide(  # noqa: PLR0913
    pres: PresentationType,
    layout: SlideLayout,
    *,
    title: str,
    text_left: str,
    text_right: str,
    highlight: str,
    picture_path: Path,
    color: RGBColor,
) -> None:
    pop_slide = pres.slides.add_slide(layout)

    title_shape = find_shape(
        pop_slide.shapes, prefix="Text Placeholder 2", kind="shape"
    )
    add_normal_text_to_shape(title_shape, title)

    text_shape = find_shape(pop_slide.shapes, prefix="Text Placeholder 3", kind="shape")
    add_highlighted_text_to_frame(
        text_shape,
        text_left=text_left,
        highlight=highlight,
        text_right=text_right,
        color=color,
    )

    figure_shape = find_shape(pop_slide.shapes, prefix="Picture", kind="picture")
    figure_shape.insert_picture(str(picture_path))


def add_built_slide(  # noqa: PLR0913
    pres: PresentationType,
    layout: SlideLayout,
    built_after_frac: float,
    *,
    start_year: int,
    pop_df: pd.DataFrame,
    built_area_df: pd.DataFrame,
    urban_area_df: pd.DataFrame,
    picture_path: Path,
) -> None:
    population = pop_df.set_index("year")["pop"] / 1e6
    built_area = built_area_df.set_index("year")["area"] / 1e6
    urban_area = urban_area_df.set_index("year")["area"] / 1e6

    pop_change = (population.loc[2020] - population.loc[start_year]) / population.loc[
        start_year
    ] + 1
    built_change = (built_area.loc[2020] - built_area.loc[start_year]) / built_area.loc[
        start_year
    ] + 1
    urban_change = (urban_area.loc[2020] - urban_area.loc[start_year]) / urban_area.loc[
        start_year
    ] + 1

    built_slide = pres.slides.add_slide(layout)

    title_shape = find_shape(
        built_slide.shapes, prefix="Text Placeholder 2", kind="shape"
    )
    add_normal_text_to_shape(
        title_shape, f"Dinámica de urbanización {start_year} a 2020"
    )

    text_shape = find_shape(
        built_slide.shapes, prefix="Text Placeholder 3", kind="shape"
    )
    add_highlighted_text_to_frame(
        text_shape,
        text_left="El ",
        highlight=f"{built_after_frac:.0%}",
        text_right=(
            f" de los píxeles se urbanizaron en el año {start_year} o posterior."
        ),
        color=RGB_BLUE,
    )

    table_shape = find_shape(built_slide.shapes, prefix="Table", kind="table")
    table_frame = table_shape.insert_table(rows=4, cols=4)
    table = table_frame.table

    table.columns[0].width = Cm(1.8)
    table.columns[1].width = Cm(3.0)
    table.columns[2].width = Cm(3.0)
    table.columns[3].width = Cm(2.9)
    table.rows[0].height = Cm(1.7)

    table_frame.element.xpath(".//a:tableStyleId")[0].text = TABLE_STYLE_ID

    text_arr = np.array(
        [
            [
                "",
                "Población\n(millones)",
                "Superficie\nconstruida\n(km²)",
                "Superficie\nurbana\n(km²)",
            ],
            [
                str(start_year),
                f"{population.loc[start_year]:.2f}",
                f"{built_area.loc[start_year]:.1f}",
                f"{urban_area.loc[start_year]:.1f}",
            ],
            [
                "2020",
                f"{population.loc[2020]:.2f}",
                f"{built_area.loc[2020]:.1f}",
                f"{urban_area.loc[2020]:.1f}",
            ],
            ["", f"x{pop_change:.2f}", f"x{built_change:.2f}", f"x{urban_change:.2f}"],
        ],
    )

    style_arr = np.array(
        [
            [False, True, True, True],
            [True, False, False, False],
            [True, False, False, False],
            [False, True, True, True],
        ],
    )

    for row_idx in range(text_arr.shape[0]):
        for col_idx in range(text_arr.shape[1]):
            cell = table.cell(row_idx, col_idx)
            cell.text = text_arr[row_idx, col_idx]
            for paragraph in cell.text_frame.paragraphs:
                if row_idx == 0:
                    paragraph.space_before = Pt(0)
                    paragraph.space_after = Pt(0)
                    paragraph.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.size = Pt(12 if row_idx == 0 else 14)
                    if style_arr[row_idx, col_idx]:
                        run.font.bold = True
                        run.font.color.rgb = RGB_BLUE

    if picture_path.exists():
        figure_shape = find_shape(built_slide.shapes, prefix="Picture", kind="picture")
        figure_shape.insert_picture(str(picture_path))


def add_single_picture_slide(
    pres: PresentationType,
    layout: SlideLayout,
    *,
    picture_path: Path,
    title: str,
) -> None:
    slide = pres.slides.add_slide(layout)

    title_shape = find_shape(slide.shapes, prefix="Text", kind="shape")
    add_normal_text_to_shape(title_shape, title)

    if picture_path.exists():
        figure_shape = find_shape(slide.shapes, prefix="Picture", kind="picture")
        figure_shape.insert_picture(str(picture_path))


def add_radial_population_slide(
    pres: PresentationType, layout: SlideLayout, comparison: RadialComparison
) -> None:
    slide = pres.slides.add_slide(layout)
    title = find_shape(slide.shapes, prefix="Text Placeholder 2", kind="shape")
    title_left, title_width = title.left, title.width
    title.left, title.width = title_left, title_width
    title.top, title.height = Cm(0.45), Cm(2.6)
    add_normal_text_to_shape(
        title,
        f"Cambio de población por distancia al centro,\n{comparison.start_year} a 2020",
    )
    for paragraph in title.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(24)
    narrative = find_shape(slide.shapes, prefix="Text Placeholder 3", kind="shape")
    narrative.text_frame.clear()

    placeholder = find_shape(slide.shapes, prefix="Picture", kind="picture")
    left, top, width, height = (
        placeholder.left,
        placeholder.top,
        placeholder.width,
        placeholder.height,
    )
    image_width, image_height = Image.from_file(str(comparison.picture_path)).size
    scale = min(width / image_width, height / image_height)
    picture = placeholder.insert_picture(str(comparison.picture_path))
    picture.width = round(image_width * scale)
    picture.height = round(image_height * scale)
    picture.left = left + (width - picture.width) // 2
    picture.top = top + (height - picture.height) // 2
    picture.crop_left = picture.crop_right = 0
    picture.crop_top = picture.crop_bottom = 0

    table_placeholder = find_shape(slide.shapes, prefix="Table", kind="table")
    table_left, table_width = table_placeholder.left, table_placeholder.width
    table_frame = table_placeholder.insert_table(rows=5, cols=3)
    table_height = table_frame.height
    table_frame.left = table_left
    table_frame.width = table_width
    table_frame.height = table_height
    table_frame.top = top + (height - table_height) // 2
    table = table_frame.table
    table_frame.element.xpath(".//a:tableStyleId")[0].text = TABLE_STYLE_ID
    table.columns[0].width = round(table_width * 0.32)
    table.columns[1].width = round(table_width * 0.38)
    table.columns[2].width = (
        table_width - table.columns[0].width - table.columns[1].width
    )
    rows = [("Zona", "Cambio de\npoblación", "Cambio\n(%)")]
    rows.extend(
        (
            str(category),
            f"{change:,.0f}",
            "N/D" if pd.isna(percent) else f"{percent:,.1f}%",
        )
        for category, change, percent in comparison.table[
            ["remoteness_cat", "pop_diff", "pop_diff_pct"]
        ].itertuples(index=False, name=None)
    )
    for row_index, row in enumerate(rows):
        for column_index, text in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.text = text
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(12 if row_index == 0 else 14)
                    if row_index == 0 or column_index == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGB_BLUE

    if comparison.table["pop_diff_pct"].isna().any():
        narrative.left, narrative.width = table_left, table_width
        narrative.top = table_frame.top + table_frame.height
        narrative.height = Cm(1)
        narrative.text = f"N/D: sin población en {comparison.start_year}."
        narrative.text_frame.paragraphs[0].runs[0].font.size = Pt(12)


def generate_single_slide(  # noqa: PLR0913
    *,
    name: str,
    comparisons: tuple[Comparison, ...],
    total_jobs: float,
    pop_df: pd.DataFrame,
    built_df: pd.DataFrame,
    built_urban_df: pd.DataFrame,
    built_figure_path: Path,
    income_figure_path: Path | None,
    jobs_figure_path: Path,
    radial_comparisons: tuple[RadialComparison, ...] = (),
) -> PresentationType:
    pres = Presentation("./template.pptx")
    layouts = find_layouts(pres)

    add_section_slide(pres, layouts["section"], name)

    for comparison in sorted(comparisons, key=lambda item: item.start_year):
        add_picture_with_highlight_slide(
            pres,
            layouts["picture_with_title_and_content"],
            title=f"Dinámica de población {comparison.start_year} a 2020",
            text_left="El ",
            highlight=f"{comparison.lost_pop_fraction:.0%}",
            text_right=(
                " de la superficie construida de 2020 perdió población"
                f" con respecto al {comparison.start_year}."
            ),
            picture_path=comparison.population_map,
            color=RGB_RED,
        )

        add_built_slide(
            pres,
            layouts["picture_with_title_content_table"],
            comparison.built_fraction,
            start_year=comparison.start_year,
            pop_df=pop_df,
            built_area_df=built_df,
            urban_area_df=built_urban_df,
            picture_path=built_figure_path,
        )

        for radial_comparison in radial_comparisons:
            if radial_comparison.start_year == comparison.start_year:
                add_radial_population_slide(
                    pres, layouts["picture_with_title_content_table"], radial_comparison
                )

    if income_figure_path is not None:
        add_single_picture_slide(
            pres,
            layouts["picture_with_title"],
            picture_path=income_figure_path,
            title="Ingreso",
        )

    add_picture_with_highlight_slide(
        pres,
        layouts["picture_with_title_and_content"],
        title="Empleos totales",
        text_left="La zona metropolitana tiene un total de ",
        highlight=f"{total_jobs:,.0f}",
        text_right=" empleos.",
        picture_path=jobs_figure_path,
        color=RGB_BLUE,
    )

    return pres


def slides_factory(
    level: str,
    *,
    partitions_def: dg.PartitionsDefinition,
    start_year: int = 2000,
) -> dg.AssetsDefinition:
    radial_ins = {}
    if level == "zone":
        name = f"radial_population_change_{start_year}_2020"
        radial_ins = {
            "radial_figure_path": dg.AssetIn(
                key=["plot_zone", name], input_manager_key="path_manager"
            ),
            "radial_table": dg.AssetIn(key=["stats_zone", name]),
        }

    @dg.asset(
        name=f"slides_zone_{start_year}_2020" if level == "zone" else f"slides_{level}",
        ins={
            **radial_ins,
            "lost_pop_fraction": dg.AssetIn(
                key=[f"stats_{level}", f"lost_pop_after_{start_year}"],
            ),
            "built_fraction": dg.AssetIn(
                key=[f"stats_{level}", f"built_after_{start_year}"]
            ),
            "pop_df": dg.AssetIn(key=[f"stats_{level}", "population"]),
            "total_jobs": dg.AssetIn(key=[f"stats_{level}", "total_jobs"]),
            "built_df": dg.AssetIn(key=[f"stats_{level}", "built_area"]),
            "built_urban_df": dg.AssetIn(key=[f"stats_{level}", "built_urban_area"]),
            "pg_figure_path": dg.AssetIn(
                key=[
                    f"plot_{level}",
                    "population_grid_1990_2020"
                    if start_year == 1990
                    else "population_grid",
                ],
                input_manager_key="path_manager",
            ),
            "built_figure_path": dg.AssetIn(
                key=[f"plot_{level}", "built"],
                input_manager_key="path_manager",
            ),
            "income_figure_path": dg.AssetIn(
                key=[f"plot_{level}", "income"],
                input_manager_key="path_manager",
            ),
            "jobs_figure_path": dg.AssetIn(
                key=[f"plot_{level}", "jobs"],
                input_manager_key="path_manager",
            ),
        },
        partitions_def=partitions_def,
        io_manager_key="presentation_manager",
        group_name=f"slides_{level}",
        required_resource_keys={f"{level}_config_resource"},
    )
    def _asset(  # noqa: PLR0913, PLR0917
        context: dg.AssetExecutionContext,
        lost_pop_fraction: float,
        built_fraction: float,
        total_jobs: float,
        pop_df: pd.DataFrame,
        built_df: pd.DataFrame,
        built_urban_df: pd.DataFrame,
        pg_figure_path: Path,
        built_figure_path: Path,
        income_figure_path: Path | None,
        jobs_figure_path: Path,
        **radial_inputs: Path | pd.DataFrame,
    ) -> PresentationType:
        config_resource = getattr(context.resources, f"{level}_config_resource")
        radial_comparisons = (
            (
                RadialComparison(
                    start_year,
                    cast("Path", radial_inputs["radial_figure_path"]),
                    cast("pd.DataFrame", radial_inputs["radial_table"]),
                ),
            )
            if level == "zone"
            else ()
        )

        return generate_single_slide(
            name=config_resource.names[context.partition_key],
            comparisons=(
                Comparison(
                    start_year, lost_pop_fraction, built_fraction, pg_figure_path
                ),
            ),
            total_jobs=total_jobs,
            pop_df=pop_df,
            built_df=built_df,
            built_urban_df=built_urban_df,
            built_figure_path=built_figure_path,
            income_figure_path=income_figure_path,
            jobs_figure_path=jobs_figure_path,
            radial_comparisons=radial_comparisons,
        )

    return _asset


slides_zone_1990_2020 = slides_factory(
    "zone", partitions_def=zone_partitions, start_year=1990
)
slides_zone_2000_2020 = slides_factory(
    "zone", partitions_def=zone_partitions, start_year=2000
)
slides_mun = slides_factory("mun", partitions_def=mun_partitions)
