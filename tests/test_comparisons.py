from pathlib import Path
from typing import cast
from unittest.mock import patch

import geopandas as gpd
import numpy as np
import pandas as pd
import pytest
from affine import Affine
from cfc_dagster_utils.resources import PostgresResource
from PIL import Image
from pptx import Presentation
from pptx.shapes.autoshape import Shape
from pptx.shapes.graphfrm import GraphicFrame
from pptx.shapes.picture import Picture
from pptx.slide import Slide
from shapely.geometry import box

import dagster as dg
from jat_slides.definitions import definitions
from jat_slides.defs.assets.agebs import agebs_factory
from jat_slides.defs.assets.cells import cells_zone_factory
from jat_slides.defs.assets.maps.population_grid import (
    population_grid_1990_2020_plot_zone,
)
from jat_slides.defs.assets.slides import (
    slides_mun,
    slides_zone_1990_2020,
    slides_zone_2000_2020,
)
from jat_slides.defs.assets.stats.built_after_2000 import built_after_factory
from jat_slides.defs.assets.stats.lost_pop_after_2000 import lost_pop_factory
from jat_slides.defs.managers import PresentationIOManager
from jat_slides.defs.partitions import mun_partitions
from jat_slides.defs.resources import ConfigResource, PathResource

ZONE = "01.1.01"


@pytest.mark.parametrize("year", [1990, 2000, 2010, 2020])
def test_census_year_query(tmp_path: Path, year: int) -> None:
    resource = PostgresResource(
        host="unused",
        port=5432,
        user="unused",
        password=str(tmp_path),
        database="unused",
    )
    frame = gpd.GeoDataFrame(
        {"pobtot": [year]}, geometry=[box(-103, 21, -102, 22)], crs=4326
    )
    with (
        patch.object(PostgresResource, "connect"),
        patch("geopandas.read_postgis", return_value=frame) as read,
        dg.build_asset_context(partition_key=ZONE) as context,
    ):
        result = agebs_factory(year)(context, resource)
    result = cast("gpd.GeoDataFrame", result)
    query = read.call_args.args[0]
    np.testing.assert_equal(read.call_args.kwargs["params"], {"zone": ZONE})
    np.testing.assert_equal(f"{year}_ageb" in query, desired=True)
    np.testing.assert_equal("pob1 AS pobtot" in query, desired=year < 2010)
    for other in {1990, 2000, 2010, 2020} - {year}:
        np.testing.assert_equal(f"{other}_ageb" in query, desired=False)
    np.testing.assert_equal(result["pobtot"].tolist(), [year])
    np.testing.assert_equal(str(result.crs), "ESRI:54009")
    np.testing.assert_equal(result.is_valid.all(), desired=True)


@pytest.mark.parametrize("year", [1990, 2000])
def test_period_cells_and_loss(tmp_path: Path, year: int) -> None:
    paths = PathResource(
        in_path=str(tmp_path),
        out_path=str(tmp_path),
        pg_path=str(tmp_path),
        job_estimation_path="unused",
    )
    for baseline, values in [(1990, [-2, -1, 0, 3]), (2000, [-1, 0, 2, 3])]:
        directory = tmp_path / "differences" / f"{baseline}_2020"
        directory.mkdir(parents=True)
        gpd.GeoDataFrame(
            {"difference": values},
            geometry=[box(i, 0, i + 1, 1) for i in range(4)],
            crs=6372,
        ).to_parquet(directory / f"{ZONE}.geoparquet")
    cells = cells_zone_factory(year)
    with dg.build_asset_context(partition_key=ZONE) as context:
        frame = cells(context, paths)
    loss = lost_pop_factory("zone", start_year=year, cells_key=cells.key)
    np.testing.assert_equal(loss(frame), 0.5 if year == 1990 else 0.25)


@pytest.mark.parametrize(("year", "expected"), [(1990, 0.75), (2000, 0.5)])
def test_construction_threshold(year: int, expected: float) -> None:
    asset = built_after_factory("mun", mun_partitions, start_year=year)
    data = np.array([[0, 1975, 1990, 2000, 2020]])
    np.testing.assert_equal(asset((data, Affine.identity())), expected)


def test_asset_dependencies() -> None:
    defs = definitions()
    dg.Definitions.validate_loadable(defs)
    graph = defs.resolve_asset_graph()
    cells_key = dg.AssetKey(["cells", "zone_1990_2020"])
    np.testing.assert_equal(
        graph.get(dg.AssetKey(["stats_zone", "lost_pop_after_1990"])).parent_keys,
        {cells_key},
    )
    np.testing.assert_equal(
        graph.get(dg.AssetKey(["stats_zone", "built_after_1990"])).parent_keys,
        {dg.AssetKey("built_zone")},
    )
    asset = population_grid_1990_2020_plot_zone
    np.testing.assert_equal(
        asset.asset_deps[asset.key],
        {cells_key, dg.AssetKey("remoteness_circles_zone")},
    )
    np.testing.assert_equal(
        dg.AssetKey("slides_zone") in graph.get_all_asset_keys(), desired=False
    )
    for year, deck in [(1990, slides_zone_1990_2020), (2000, slides_zone_2000_2020)]:
        other_year = 2000 if year == 1990 else 1990
        np.testing.assert_equal(deck.key.path, [f"slides_zone_{year}_2020"])
        np.testing.assert_equal(
            deck.partitions_def, population_grid_1990_2020_plot_zone.partitions_def
        )
        parents = deck.asset_deps[deck.key]
        for metric in ("lost_pop_after", "built_after"):
            np.testing.assert_equal(
                dg.AssetKey(["stats_zone", f"{metric}_{year}"]) in parents, desired=True
            )
            np.testing.assert_equal(
                dg.AssetKey(["stats_zone", f"{metric}_{other_year}"]) in parents,
                desired=False,
            )
        np.testing.assert_equal(
            deck.keys_by_input_name["pg_figure_path"].path,
            [
                "plot_zone",
                "population_grid_1990_2020" if year == 1990 else "population_grid",
            ],
        )
        for prefix in ("plot_zone", "stats_zone"):
            np.testing.assert_equal(
                dg.AssetKey([prefix, f"radial_population_change_{other_year}_2020"])
                in parents,
                desired=False,
            )
        for name in ("income", "jobs", "built"):
            np.testing.assert_equal(
                dg.AssetKey(["plot_zone", name]) in parents, desired=True
            )
        np.testing.assert_equal(
            deck.get_io_manager_key_for_asset_key(deck.key), "presentation_manager"
        )
        np.testing.assert_equal(
            deck.op.input_def_named("pg_figure_path").input_manager_key, "path_manager"
        )
    np.testing.assert_equal(slides_mun.key.path, ["slides_mun"])


def slide_text(slide: Slide) -> str:
    return " ".join(shape.text for shape in slide.shapes if isinstance(shape, Shape))


def slide_image(slide: Slide) -> bytes:
    return next(
        shape.image.blob for shape in slide.shapes if isinstance(shape, Picture)
    )


def source_asset(
    key: dg.AssetKey, partitions: dg.PartitionsDefinition, value: object
) -> dg.AssetsDefinition:
    @dg.asset(key=key, partitions_def=partitions)
    def source() -> object:
        return value

    return source


@pytest.mark.parametrize(
    ("level", "year"), [("zone", 1990), ("zone", 2000), ("mun", 2000)]
)
@pytest.mark.parametrize("with_income", [True, False])
def test_slideshow(tmp_path: Path, level: str, year: int, *, with_income: bool) -> None:
    images = {}
    for name, color in [("built", "blue"), ("2000", "red"), ("1990", "green")]:
        path = tmp_path / f"{name}.png"
        Image.new("RGB", (100, 100), color).save(path)
        images[name] = path
    asset = {
        ("zone", 1990): slides_zone_1990_2020,
        ("zone", 2000): slides_zone_2000_2020,
        ("mun", 2000): slides_mun,
    }[level, year]
    inputs = {
        "lost_pop_fraction": 0.25 if year == 1990 else 0.5,
        "built_fraction": 0.75 if year == 1990 else 0.25,
        "total_jobs": 1234.0,
        "pop_df": pd.DataFrame({"year": [1990, 2000, 2020], "pop": [1e6, 2e6, 4e6]}),
        "built_df": pd.DataFrame({"year": [1990, 2000, 2020], "area": [2e6, 4e6, 8e6]}),
        "built_urban_df": pd.DataFrame(
            {"year": [1990, 2000, 2020], "area": [3e6, 6e6, 12e6]}
        ),
        "built_figure_path": images["built"],
        "pg_figure_path": images[str(year)],
        "income_figure_path": images["built"] if with_income else None,
        "jobs_figure_path": images["built"],
    }
    if level == "zone":
        inputs["radial_figure_path"] = images[str(year)]
        inputs["radial_table"] = pd.DataFrame(
            {
                "remoteness_cat": ["Central", "Intermedia", "Distante", "Periurbana"],
                "pop_diff": [-1234.5, 2000, 0, year],
                "pop_diff_pct": [-12.345, 200, 0, np.nan],
            }
        )
    partition = ZONE if level == "zone" else "01001"
    paths = PathResource(
        in_path=str(tmp_path),
        out_path=str(tmp_path),
        pg_path="unused",
        job_estimation_path="unused",
    )
    memory = dg.InMemoryIOManager()
    partitions = asset.partitions_def
    if partitions is None:
        pytest.fail("Slides must remain partitioned")
    result = dg.materialize(
        [
            asset,
            *[
                source_asset(key, partitions, inputs[name])
                for name, key in asset.keys_by_input_name.items()
            ],
        ],
        partition_key=partition,
        resources={
            "io_manager": memory,
            "path_manager": memory,
            "presentation_manager": PresentationIOManager(
                path_resource=paths, extension=".pptx"
            ),
            f"{level}_config_resource": ConfigResource(
                bounds={}, names={partition: "Aguascalientes"}
            ),
        },
    )
    np.testing.assert_equal(result.success, desired=True)
    name = f"slides_zone_{year}_2020" if level == "zone" else "slides_mun"
    slides = Presentation(str(tmp_path / name / f"{partition}.pptx")).slides
    np.testing.assert_equal(len(slides), 4 + int(with_income) + (level == "zone"))
    np.testing.assert_equal(slide_text(slides[0]), "Aguascalientes")
    np.testing.assert_equal(
        f"Dinámica de población {year} a 2020" in slide_text(slides[1]), desired=True
    )
    np.testing.assert_equal(
        ("25%" if year == 1990 else "50%") in slide_text(slides[1]), desired=True
    )
    np.testing.assert_equal(slide_image(slides[1]), images[str(year)].read_bytes())
    np.testing.assert_equal(
        f"Dinámica de urbanización {year} a 2020" in slide_text(slides[2]), desired=True
    )
    np.testing.assert_equal(
        ("75%" if year == 1990 else "25%") in slide_text(slides[2]), desired=True
    )
    table = next(shape.table for shape in slides[2].shapes if shape.has_table)
    np.testing.assert_equal(table.cell(1, 0).text, str(year))
    np.testing.assert_equal(table.cell(2, 0).text, "2020")
    np.testing.assert_equal(
        [table.cell(1, col).text for col in range(1, 4)],
        ["1.00", "2.0", "3.0"] if year == 1990 else ["2.00", "4.0", "6.0"],
    )
    np.testing.assert_equal(
        [table.cell(2, col).text for col in range(1, 4)], ["4.00", "8.0", "12.0"]
    )
    np.testing.assert_equal(
        [table.cell(3, col).text for col in range(1, 4)],
        ["x4.00" if year == 1990 else "x2.00"] * 3,
    )
    np.testing.assert_equal(slide_image(slides[2]), images["built"].read_bytes())
    other_year = 2000 if year == 1990 else 1990
    for slide in slides:
        np.testing.assert_equal(
            f"{other_year} a 2020" in slide_text(slide), desired=False
        )
    if with_income:
        np.testing.assert_equal(slide_text(slides[-2]), "Ingreso")
        np.testing.assert_equal(slide_image(slides[-2]), images["built"].read_bytes())
    np.testing.assert_equal("Empleos totales" in slide_text(slides[-1]), desired=True)
    np.testing.assert_equal("1,234" in slide_text(slides[-1]), desired=True)
    np.testing.assert_equal(slide_image(slides[-1]), images["built"].read_bytes())
    if level == "zone":
        assert_radial_slide(slides[3], year, images[str(year)])


def assert_radial_slide(slide: Slide, year: int, picture_path: Path) -> None:
    np.testing.assert_equal(
        f"Cambio de población por distancia al centro, {year} a 2020"
        in slide_text(slide).replace("\n", " "),
        desired=True,
    )
    picture = next(s for s in slide.shapes if isinstance(s, Picture))
    frame = next(s for s in slide.shapes if isinstance(s, GraphicFrame) and s.has_table)
    np.testing.assert_equal(picture.image.blob, picture_path.read_bytes())
    np.testing.assert_equal(
        [
            picture.crop_left,
            picture.crop_right,
            picture.crop_top,
            picture.crop_bottom,
        ],
        [0, 0, 0, 0],
    )
    np.testing.assert_equal(picture.left + picture.width <= frame.left, desired=True)
    table = frame.table
    np.testing.assert_equal(len(table.rows), 5)
    np.testing.assert_equal(len(table.columns), 3)
    np.testing.assert_equal(table.cell(0, 1).text, "Cambio de\npoblación")
    np.testing.assert_equal(table.cell(1, 1).text, "-1,234")
    np.testing.assert_equal(table.cell(4, 1).text, f"{year:,}")
    np.testing.assert_equal(table.cell(0, 2).text, "Cambio\n(%)")
    np.testing.assert_equal(
        [table.cell(row, 2).text for row in range(1, 5)],
        ["-12.3%", "200.0%", "0.0%", "N/D"],
    )
    np.testing.assert_equal(
        f"N/D: sin población en {year}." in slide_text(slide), desired=True
    )
